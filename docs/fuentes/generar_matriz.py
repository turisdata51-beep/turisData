"""Genera docs/fuentes/matriz_indicadores_fuentes.xlsx.

Cruza cada indicador del Manual con sus variables, la fuente que puede darlas y el
método para obtenerlas (scraping/descarga automática, solicitud a la Cámara,
encuesta propia, solicitud a otra entidad o no factible).
Fuentes verificadas el 27 de octubre de 2026.

Uso: python generar_matriz.py   (requiere openpyxl)
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from indicadores_variables import INDICADORES

CARPETA = Path(__file__).resolve().parent

AUTO = "Scraping / descarga automática"
CAMARA = "Solicitar a la Cámara"
ENCUESTA = "Encuesta propia"
OTRA = "Solicitar a otra entidad"
NO = "No factible por ahora"
INTERNO = "Cálculo interno del sistema"

COLORES = {AUTO: "D9EAD3", CAMARA: "FCE5CD", ENCUESTA: "FFF2CC", OTRA: "F4CCCC", NO: "E6E6E6", INTERNO: "DDEBF7"}

# código -> (método principal, fuente principal, complemento o alternativa, nivel geográfico, frecuencia, observación)
CLASIFICACION = {
    "A1": (CAMARA, "Reporte mensual de los alojamientos asociados (plantilla v1)", "Informes de ocupación de feriados de la Fundación Turismo para Cuenca (PDF, esporádico)", "Cantón / parroquia", "Mensual", "No existe fuente pública con llegadas mensuales a nivel Cuenca."),
    "A2": (CAMARA, "Reporte mensual de los alojamientos asociados", "", "Cantón / parroquia", "Mensual", "Base de A4, A5, B5, D4 y E3–E5."),
    "A3": (ENCUESTA, "Conteos o encuesta en puntos de interés", "Visitantes por feriado en notas de prensa del GAD Municipal (scraping de noticias, sin serie tabulada)", "Cantón", "Por campaña / feriado", "Las notas municipales mezclan turistas y excursionistas."),
    "A4": (INTERNO, "Se calcula con A2 / A1", "", "Cantón / parroquia", "Mensual", "Depende de los datos de la Cámara."),
    "A5": (CAMARA, "Reporte mensual (noches de residentes y no residentes)", "", "Cantón", "Mensual", "La plantilla ya separa residentes y no residentes."),
    "A6": (ENCUESTA, "Encuesta de gasto a turistas", "Gasto promedio en informes de feriados del GAD Municipal", "Cantón", "Semestral / feriado", "Requiere diseño muestral."),
    "A7": (ENCUESTA, "Encuesta de gasto a excursionistas", "", "Cantón", "Semestral / feriado", ""),
    "A8": (ENCUESTA, "Encuesta a visitantes (pregunta de visita previa en 5 años)", "", "Cantón", "Semestral", ""),
    "A9": (ENCUESTA, "Encuesta a visitantes (satisfacción global)", "", "Cantón", "Semestral", "Definir escala y umbral antes de iniciar."),
    "B1": (AUTO, "Catastro Turístico, Portal de Datos Abiertos del Ecuador (XLSX vía API CKAN)", "Hostelería, Cuenca en Datos (CSV vía API CKAN) y catastro de los asociados de la Cámara", "Cantón / parroquia", "Irregular (último: feb-2025 y abr-2025)", "1.599 registros de Cuenca en el catastro nacional. Trae datos personales: se guardan solo agregados."),
    "B2": (AUTO, "Habitaciones instaladas: Hostelería, Cuenca en Datos (nro_habitaciones)", "Habitaciones-noche disponibles: reporte mensual de la Cámara", "Cantón / parroquia", "Irregular / mensual", "La capacidad instalada se automatiza; la vendible mes a mes viene de la Cámara."),
    "B3": (AUTO, "Plazas instaladas: Hostelería, Cuenca en Datos (nro_plazas)", "Plazas-noche disponibles: reporte mensual de la Cámara", "Cantón / parroquia", "Irregular / mensual", ""),
    "B4": (CAMARA, "Reporte mensual de la Cámara", "Ocupación de feriados en informes de la Fundación Turismo (PDF) como referencia", "Cantón / parroquia", "Mensual", "La Encuesta Coyuntural de Alojamiento no se publica a nivel Cuenca (no verificado en el portal caído)."),
    "B5": (CAMARA, "Reporte mensual de la Cámara", "", "Cantón", "Mensual", ""),
    "B6": (CAMARA, "Reporte mensual (ingreso de habitaciones)", "", "Cantón", "Mensual", "Confirmar que los hoteles compartan ingresos."),
    "B7": (CAMARA, "Reporte mensual (ingreso de habitaciones)", "", "Cantón", "Mensual", ""),
    "B8": (CAMARA, "Ingreso total de los asociados (confidencial, agregado)", "", "Cantón", "Mensual", "Fase avanzada."),
    "B9": (CAMARA, "GOP de los asociados (confidencial, agregado)", "", "Cantón", "Mensual", "Fase avanzada; solo con mínimo de participantes."),
    "C1": (NO, "Airbnb / Booking (prohibido el scraping en sus términos de uso)", "AirDNA con licencia de pago o académica", "Cantón", "Mensual", "Airbnb sec. 11.1 y Booking A15.2 prohíben bots y scrapers. Inside Airbnb no cubre Ecuador."),
    "C2": (NO, "Igual que C1", "AirDNA con licencia", "Cantón", "Mensual", ""),
    "C3": (NO, "Igual que C1", "AirDNA con licencia", "Cantón", "Mensual", ""),
    "C4": (NO, "Igual que C1", "AirDNA con licencia", "Cantón", "Mensual", ""),
    "C5": (NO, "Igual que C1", "AirDNA con licencia", "Cantón", "Mensual", ""),
    "C6": (NO, "Igual que C1", "AirDNA con licencia", "Cantón", "Mensual", ""),
    "C7": (CAMARA, "Fecha de reserva y de llegada de los asociados (si su sistema lo registra)", "AirDNA con licencia para STR", "Cantón", "Mensual", "Se puede aplicar a hoteles si exportan reservas."),
    "C8": (CAMARA, "Reservas futuras de los asociados (fotografías periódicas)", "", "Cantón", "Semanal", "Requiere integración con el sistema de reservas de cada hotel."),
    "D1": (AUTO, "Ventas de la sección CIIU I en Cuenca, SRI datos abiertos (scraping de sri.gob.ec/datasets) — YA IMPLEMENTADO", "Impacto económico por feriado (GAD Municipal) y encuesta de gasto", "Cantón", "Mensual (hasta ago-2026)", "Es un indicador de actividad (ventas de alojamiento + comidas), no el gasto turístico exacto. 104 meses ya cargados (2018–2026)."),
    "D2": (AUTO, "ENEMDU, INEC (microdatos CSV; Cuenca es ciudad autorrepresentada)", "Empleo de los asociados de la Cámara", "Ciudad", "Trimestral", "Muestra pequeña para la rama I: error alto. Publicar con advertencia."),
    "D3": (CAMARA, "Empleo estacional de los asociados (encuesta empresarial)", "", "Cantón", "Trimestral", ""),
    "D4": (AUTO, "Población: proyecciones INEC Revisión 2024 (XLSX)", "Noches: reporte de la Cámara", "Cantón", "Mensual", "Población 2026 proyectada: 628.740; censo 2022: 596.101."),
    "D5": (ENCUESTA, "Visitantes totales (estimación propia o municipal)", "Población INEC (automática)", "Cantón", "Feriado / anual", "Sin una fuente de visitantes totales solo se puede para feriados."),
    "D6": (AUTO, "Plazas: Hostelería, Cuenca en Datos + población INEC", "", "Cantón / parroquia", "Anual", "Totalmente automatizable."),
    "D7": (ENCUESTA, "Encuesta a residentes", "", "Cantón", "Anual", ""),
    "D8": (NO, "Cuenta Satélite de Turismo (solo nacional, BCE)", "", "Nacional", "Anual", "No existe estimación cantonal."),
    "E1": (CAMARA, "Auditoría de accesibilidad de los asociados", "", "Cantón", "Anual", ""),
    "E2": (ENCUESTA, "Encuesta a visitantes (modo de transporte)", "", "Cantón", "Semestral", ""),
    "E3": (OTRA, "Consumo de agua de los hoteles (ETAPA EP o facturas vía Cámara)", "", "Cantón", "Mensual", "ETAPA no publica consumo por sector."),
    "E4": (OTRA, "Consumo de energía de los hoteles (Centrosur o facturas vía Cámara)", "", "Cantón", "Mensual", "Centrosur publica por área de concesión, sin sector turístico."),
    "E5": (OTRA, "Residuos de los hoteles (EMAC o pesaje propio)", "", "Cantón", "Mensual", "EMAC no publica por actividad."),
    "E6": (CAMARA, "Certificaciones vigentes de los asociados", "Distintivos publicados por el Viceministerio de Turismo (no verificado)", "Cantón", "Anual", ""),
    "Q1": (INTERNO, "Registro de cargas del sistema", "", "Cantón", "Mensual", ""),
    "Q2": (INTERNO, "Validador de cargas", "", "—", "Por carga", ""),
    "Q3": (INTERNO, "Bitácora de cargas y de extractores", "", "—", "Diaria", ""),
}

FUENTES = [
    # (fuente, URL, qué ofrece, nivel, periodicidad / último dato, formato, método, restricciones, dificultad, estado)
    ("SRI – Ventas por cantón y actividad", "https://www.sri.gob.ec/datasets", "Ventas mensuales por cantón y sección CIIU (I = alojamiento y comidas)", "Cantón", "Mensual / ago-2026", "CSV (formato varía por año)", "Scraping de la página + descarga", "robots.txt permite todo", "Baja", "Implementado"),
    ("SRI – Recaudación", "https://www.sri.gob.ec/datasets", "Recaudación por impuesto, cantón y sección CIIU", "Cantón", "Mensual / 2026", "CSV", "Scraping + descarga", "robots.txt permite todo", "Baja", "Pendiente"),
    ("Datos Abiertos Ecuador – Catastro Turístico", "https://www.datosabiertos.gob.ec", "Establecimientos turísticos con actividad, categoría, cantón y parroquia", "Establecimiento", "Irregular / feb-2025", "XLSX vía API CKAN", "Descarga automática por API", "CC-BY; contiene datos personales", "Baja", "Pendiente"),
    ("Cuenca en Datos – Hostelería", "https://cuencaendatos.cuenca.gob.ec/dataset/hosteleria", "Alojamientos con tipo, categoría, habitaciones y plazas", "Establecimiento", "Irregular / abr-2025", "CSV vía API CKAN", "Descarga automática por API", "CC No Comercial", "Baja", "Pendiente"),
    ("Cuenca en Datos – Equipamiento turístico", "https://cuencaendatos.cuenca.gob.ec", "Equipamiento turístico con coordenadas", "Punto", "Irregular / abr-2025", "CSV vía API CKAN", "Descarga automática por API", "CC No Comercial", "Baja", "Pendiente"),
    ("INEC – Proyecciones de población", "https://www.ecuadorencifras.gob.ec", "Población cantonal 2010–2035", "Cantón", "Revisión 2024", "XLSX", "Descarga automática", "robots.txt solo bloquea /wp-admin/", "Baja", "Pendiente"),
    ("INEC – ENEMDU", "https://www.ecuadorencifras.gob.ec/enemdu-trimestral/", "Empleo por rama; Cuenca ciudad autorrepresentada", "Ciudad", "Trimestral / I-2026", "CSV microdatos", "Descarga automática + cálculo ponderado", "Libre", "Media", "Pendiente"),
    ("MPCEI – Boletín de cifras del sector turismo", "https://www.produccion.gob.ec/boletin-de-cifras-sector-turismo/", "Llegadas internacionales, feriados, catastro por provincia, Cajas, empleo", "Nacional / provincial", "Mensual / sep-2026", "PDF", "Scraping de la página + lectura de PDF", "Citar fuente", "Media", "Pendiente (contexto provincial)"),
    ("Fundación Turismo para Cuenca – Estadísticas", "https://cuenca.com.ec/estadisticas/", "Informes de ocupación hotelera en feriados", "Cantón", "Esporádico / Carnaval 2026", "PDF", "Scraping + lectura de PDF", "Usar user-agent identificado y ritmo bajo", "Media", "Pendiente"),
    ("GAD Municipal de Cuenca – Noticias de feriados", "https://www.cuenca.gob.ec", "Visitantes, impacto económico y ocupación por feriado", "Cantón", "Por feriado", "Notas de prensa HTML", "Scraping de noticias + revisión manual", "Sin serie tabulada", "Alta", "Pendiente"),
    ("Aeropuerto Mariscal Lamar (CORPAC)", "https://aeropuertocuenca.ec", "Pasajeros anuales (531.629 en 2025)", "Cantón", "Anual", "PDF rendición de cuentas", "Lectura de PDF; serie mensual por solicitud", "", "Alta", "Solicitar serie mensual"),
    ("Terminal Terrestre de Cuenca", "—", "Llegadas y salidas de pasajeros", "Cantón", "—", "Solo prensa", "Solicitar a la entidad", "", "Alta", "Solicitar"),
    ("ETAPA EP – Parque Nacional Cajas", "https://www.etapa.net.ec", "Visitantes del Cajas", "Área protegida", "Semestral en boletín MPCEI", "PDF / solicitud", "Solicitar serie mensual", "", "Alta", "Solicitar"),
    ("BCE – Balanza de pagos (cuenta Viajes)", "https://contenido.bce.fin.ec", "Ingresos por turismo receptor", "Nacional", "Trimestral / II-2026", "PDF", "Lectura de PDF", "", "Media", "Contexto nacional"),
    ("Booking / Airbnb / TripAdvisor / Google Maps", "—", "Precios, reseñas, anuncios", "—", "—", "—", "NO PERMITIDO", "Sus términos prohíben scraping", "—", "Descartado"),
    ("AirDNA", "https://www.airdna.co", "Indicadores de alquiler de corta estancia", "Cantón", "Mensual", "Plataforma de pago", "Licencia", "De pago", "—", "Evaluar licencia"),
    ("Cámara de Turismo del Azuay", "—", "Reporte mensual de alojamientos asociados (plantilla v1)", "Establecimiento", "Mensual", "XLSX / CSV", "Carga en el panel admin", "Convenio y confidencialidad", "Baja", "Implementado"),
]

SOLICITUD_CAMARA = [
    # (dato, para qué indicadores, frecuencia, formato)
    ("Reporte mensual de cada alojamiento asociado: habitaciones y plazas disponibles, habitaciones ocupadas, llegadas y noches por residencia", "A1, A2, A4, A5, B2–B5, D4", "Mensual (hasta el día 15)", "Plantilla v1 (XLSX/CSV) en el panel"),
    ("Ingreso de habitaciones del mes (sin IVA ni servicio)", "B6, B7", "Mensual", "Columna opcional de la plantilla v1"),
    ("Padrón de asociados: código, tipo, categoría, parroquia y fecha de participación", "B1, Q1", "Al inicio y con cada cambio", "Panel admin – Establecimientos"),
    ("Empleo de los asociados: total y estacional", "D2, D3", "Trimestral", "Nueva plantilla"),
    ("Habitaciones accesibles y certificaciones vigentes", "E1, E6", "Anual", "Nueva plantilla"),
    ("Consumo de agua, energía y residuos de hoteles voluntarios", "E3, E4, E5", "Mensual", "Nueva plantilla (fase avanzada)"),
    ("Apoyo para aplicar encuestas a visitantes y residentes", "A3, A6–A9, D1, D5, D7, E2", "Semestral / feriados", "Formulario en línea"),
    ("Gestionar con CORPAC, la Terminal y ETAPA las series mensuales de pasajeros y visitantes del Cajas", "Contexto de demanda", "Mensual", "Oficio institucional"),
]


def hoja(wb, titulo, encabezados, filas, anchos, color_col=None):
    ws = wb.create_sheet(titulo)
    ws.append(encabezados)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="A5162D")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for f in filas:
        ws.append(list(f))
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if color_col is not None:
            metodo = fila[color_col].value
            if metodo in COLORES:
                fila[color_col].fill = PatternFill("solid", fgColor=COLORES[metodo])
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def main():
    wb = Workbook()
    resumen = wb.active
    resumen.title = "Resumen"

    filas = []
    for codigo, nombre, prioridad, fase, variables, calculo in INDICADORES:
        metodo, fuente, alterna, nivel, frecuencia, obs = CLASIFICACION[codigo]
        filas.append((codigo, nombre, prioridad, fase, ", ".join(variables), calculo, metodo, fuente, alterna, nivel, frecuencia, obs))

    conteo = {}
    for f in filas:
        conteo[f[6]] = conteo.get(f[6], 0) + 1
    resumen.append(["Matriz de indicadores y fuentes – TurisData (Observatorio Turístico de Cuenca)"])
    resumen["A1"].font = Font(bold=True, size=14)
    resumen.append(["Fuentes verificadas el 27 de octubre de 2026. El Ministerio de Turismo fue fusionado al MPCEI y su portal 'Turismo en cifras' no está en línea."])
    resumen.append([])
    resumen.append(["Método de obtención", "Indicadores", "Códigos"])
    for c in resumen[4]:
        c.font = Font(bold=True)
    for metodo in (AUTO, CAMARA, ENCUESTA, OTRA, NO, INTERNO):
        codigos = [f[0] for f in filas if f[6] == metodo]
        resumen.append([metodo, len(codigos), ", ".join(codigos)])
        resumen.cell(row=resumen.max_row, column=1).fill = PatternFill("solid", fgColor=COLORES[metodo])
    resumen.column_dimensions["A"].width = 34
    resumen.column_dimensions["B"].width = 12
    resumen.column_dimensions["C"].width = 70

    hoja(wb, "Indicadores", ["Código", "Indicador", "Prioridad", "Fase", "Variables necesarias", "Cálculo", "Método",
                              "Fuente principal", "Complemento / alternativa", "Nivel geográfico", "Frecuencia", "Observaciones"],
         filas, [8, 30, 10, 10, 34, 28, 26, 46, 40, 16, 18, 46], color_col=6)
    hoja(wb, "Fuentes", ["Fuente", "URL", "Qué ofrece", "Nivel geográfico", "Periodicidad / último dato", "Formato",
                         "Método de extracción", "Restricciones", "Dificultad", "Estado en el sistema"],
         FUENTES, [36, 40, 44, 16, 22, 20, 30, 30, 11, 22])
    hoja(wb, "Solicitud a la Cámara", ["Dato a solicitar", "Indicadores", "Frecuencia", "Formato"],
         SOLICITUD_CAMARA, [70, 26, 22, 34])

    salida = CARPETA / "matriz_indicadores_fuentes.xlsx"
    wb.save(salida)
    print(f"Generado {salida.name}: {len(filas)} indicadores, {len(FUENTES)} fuentes")
    for metodo, n in conteo.items():
        print(f"  {metodo}: {n}")


if __name__ == "__main__":
    main()
