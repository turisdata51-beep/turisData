"""Plantillas complementarias: empleo, accesibilidad, consumos y series de otras entidades.

También verifica que la plantilla Excel que se descarga coincide con lo que el validador espera.
"""
import csv
import io
from datetime import date

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from app.nucleo.db import transaccion
from app.servicios.plantillas import PLANTILLAS, importar
from conftest import csv_reporte, fila_reporte, observacion, usuario


def archivo(plantilla: str, *filas: dict) -> bytes:
    columnas = [c.nombre for c in PLANTILLAS[plantilla].columnas]
    salida = io.StringIO()
    escritor = csv.DictWriter(salida, fieldnames=columnas)
    escritor.writeheader()
    for f in filas:
        escritor.writerow({c: f.get(c, "") for c in columnas})
    return salida.getvalue().encode("utf-8")


def errores(resultado):
    return {(e["campo"], e["mensaje"]) for e in resultado["errores"] if e["severidad"] == "error"}


def contar(tabla: str, vigente=True) -> int:
    with transaccion() as con:
        return con.execute(text(f"SELECT count(*) FROM {tabla} WHERE vigente = :v"), dict(v=vigente)).scalar()


# --- Empleo -------------------------------------------------------------------------------------

def test_empleo_valida_y_calcula_d3():
    r = importar(archivo("empleo",
        dict(trimestre="2026-T3", codigo_establecimiento="EST-01", empleo_total=20, empleo_estacional=5),
        dict(trimestre="2026-T3", codigo_establecimiento="EST-02", empleo_total=10, empleo_estacional=5, empleo_mujeres=6),
        dict(trimestre="2026-T3", codigo_establecimiento="EST-03", empleo_total=10, empleo_estacional=0),
        dict(trimestre="2026-T3", codigo_establecimiento="EST-04", empleo_total=5, empleo_estacional=8),   # inválida
        dict(trimestre="2026-T5", codigo_establecimiento="EST-04", empleo_total=5, empleo_estacional=1),   # inválida
    ), "empleo.csv", usuario("analista"), "empleo")
    assert (r["estado"], r["registros_validos"], r["plantilla"]) == ("con_errores", 3, "empleo")
    assert ("empleo_estacional", "El empleo estacional supera el empleo total") in errores(r)
    assert ("trimestre", "Use el formato AAAA-T1 a AAAA-T4") in errores(r)
    d3 = observacion("D3", date(2026, 7, 1))
    assert float(d3["valor"]) == pytest.approx(10 / 40 * 100, abs=1e-4)
    assert (d3["periodicidad"], d3["periodo_fin"], d3["publicable"]) == ("trimestral", date(2026, 9, 30), True)


def test_empleo_reenvio_reemplaza():
    fila = dict(trimestre="2026-T1", codigo_establecimiento="EST-01", empleo_total=20, empleo_estacional=2)
    importar(archivo("empleo", fila), "a.csv", usuario("analista"), "empleo")
    importar(archivo("empleo", fila | {"empleo_total": 22}), "b.csv", usuario("analista"), "empleo")
    assert contar("operativo.reporte_empleo_trimestral") == 1
    assert contar("operativo.reporte_empleo_trimestral", vigente=False) == 1


# --- Accesibilidad y certificaciones --------------------------------------------------------------

def test_accesibilidad_calcula_e1_y_e6():
    filas = [
        dict(anio=2026, codigo_establecimiento="EST-01", habitaciones_auditadas=20, habitaciones_accesibles=2,
             certificacion="Calidad", certificacion_vigente_hasta="2027-03-01"),
        dict(anio=2026, codigo_establecimiento="EST-02", habitaciones_auditadas=20, habitaciones_accesibles=0,
             certificacion="Calidad", certificacion_vigente_hasta="2025-06-30"),        # venció antes del año
        dict(anio=2026, codigo_establecimiento="EST-03", habitaciones_auditadas=10, habitaciones_accesibles=1),
    ]
    r = importar(archivo("accesibilidad", *filas), "acc.csv", usuario("analista"), "accesibilidad")
    assert r["estado"] == "publicada"
    e1, e6 = observacion("E1", date(2026, 1, 1)), observacion("E6", date(2026, 1, 1))
    assert float(e1["valor"]) == pytest.approx(3 / 50 * 100, abs=1e-4)
    assert float(e6["valor"]) == pytest.approx(1 / 3 * 100, abs=1e-4)
    assert (e1["periodicidad"], e1["periodo_fin"]) == ("anual", date(2026, 12, 31))


def test_accesibilidad_reglas():
    r = importar(archivo("accesibilidad",
        dict(anio=2026, codigo_establecimiento="EST-01", habitaciones_auditadas=5, habitaciones_accesibles=6),
        dict(anio=2026, codigo_establecimiento="EST-02", habitaciones_auditadas=5, habitaciones_accesibles=1, certificacion="Calidad"),
        dict(anio=26, codigo_establecimiento="EST-03", habitaciones_auditadas=5, habitaciones_accesibles=1),
    ), "acc.csv", usuario("analista"), "accesibilidad")
    assert r["estado"] == "rechazada"
    assert {"habitaciones_accesibles", "certificacion_vigente_hasta", "anio"} <= {c for c, _ in errores(r)}


# --- Consumos -------------------------------------------------------------------------------------

def test_consumos_por_noche_solo_con_las_noches_de_quienes_reportan():
    importar(csv_reporte(*[fila_reporte(c) for c in ("EST-01", "EST-02", "EST-03")]), "aloj.csv", usuario("analista"))
    r = importar(archivo("consumos",
        dict(periodo="2026-09", codigo_establecimiento="EST-01", agua_m3=90, energia_kwh=3000, residuos_kg=270),
        dict(periodo="2026-09", codigo_establecimiento="EST-02", agua_m3=60, energia_kwh=3600),
        dict(periodo="2026-09", codigo_establecimiento="EST-03", agua_m3=75, residuos_kg=300),
    ), "consumos.csv", usuario("analista"), "consumos")
    assert r["estado"] == "publicada"
    septiembre = date(2026, 9, 1)  # cada hotel reporta 300 noches en fila_reporte
    assert float(observacion("E3", septiembre)["valor"]) == pytest.approx(225 * 1000 / 900, abs=1e-4)
    assert float(observacion("E4", septiembre)["valor"]) == pytest.approx(6600 / 600, abs=1e-4)
    e5 = observacion("E5", septiembre)
    assert float(e5["valor"]) == pytest.approx(570 / 600, abs=1e-4)
    assert e5["n_unidades"] == 2 and not e5["publicable"]       # solo 2 hoteles reportan residuos


def test_consumos_requiere_al_menos_un_valor():
    r = importar(archivo("consumos", dict(periodo="2026-09", codigo_establecimiento="EST-01")),
                 "vacio.csv", usuario("analista"), "consumos")
    assert ("agua_m3", "Reporte al menos uno de los consumos") in errores(r)


def test_asociado_solo_su_establecimiento_en_cualquier_plantilla():
    r = importar(archivo("consumos", dict(periodo="2026-09", codigo_establecimiento="EST-02", agua_m3=5)),
                 "otro.csv", usuario("asociado"), "consumos")
    assert ("codigo_establecimiento", "Solo puede reportar su propio establecimiento") in errores(r)


# --- Series de otras entidades ---------------------------------------------------------------------

def serie(**cambios):
    return dict(variable="aeropuerto_pasajeros_llegada", zona="Cuenca", periodicidad="mensual",
                periodo_inicio="2026-09-01", periodo_fin="2026-09-30", valor=21000) | cambios


def test_series_externas_se_guardan_con_historial():
    importar(archivo("series_externas", serie(), serie(variable="feriado_visitantes", periodicidad="feriado",
             periodo_inicio="2026-02-14", periodo_fin="2026-02-17", evento="Carnaval 2026", valor=121411)),
             "series.csv", usuario("analista"), "series_externas")
    r = importar(archivo("series_externas", serie(valor=22500)), "correccion.csv", usuario("analista"), "series_externas")
    assert r["estado"] == "publicada"
    with transaccion() as con:
        filas = con.execute(text("""SELECT variable, evento, valor, vigente, carga_id IS NOT NULL AS con_carga
                                    FROM operativo.observacion_externa ORDER BY id""")).mappings().all()
    aeropuerto = [(float(f["valor"]), f["vigente"]) for f in filas if f["variable"] == "aeropuerto_pasajeros_llegada"]
    assert aeropuerto == [(21000, False), (22500, True)]
    assert any(f["evento"] == "Carnaval 2026" for f in filas) and all(f["con_carga"] for f in filas)


def test_series_externas_reglas():
    r = importar(archivo("series_externas",
        serie(variable="no_existe"),
        serie(zona="Atlántida"),
        serie(periodicidad="feriado"),                                   # sin nombre del feriado
        serie(periodo_fin="2026-08-01"),                                 # termina antes de empezar
    ), "malo.csv", usuario("analista"), "series_externas")
    assert r["estado"] == "rechazada"
    assert {"variable", "zona", "evento", "periodo_fin"} == {c for c, _ in errores(r)}


# --- Plantillas descargables -----------------------------------------------------------------------

@pytest.mark.parametrize("codigo", list(PLANTILLAS))
def test_excel_generado_tiene_las_columnas_del_validador(codigo):
    plantilla = PLANTILLAS[codigo]
    libro = load_workbook(io.BytesIO(plantilla.generar_xlsx()))
    assert libro.sheetnames == ["Instrucciones", "Datos", "Listas", "Diccionario"]
    encabezado = [c.value for c in libro["Datos"][1]]
    assert encabezado == [c.nombre for c in plantilla.columnas]


def test_excel_de_series_lista_solo_variables_manuales():
    libro = load_workbook(io.BytesIO(PLANTILLAS["series_externas"].generar_xlsx()))
    columna = next(col for col in libro["Listas"].iter_cols(values_only=True) if col[0] == "variable")
    variables = [v for v in columna[1:] if v]
    assert "aeropuerto_pasajeros_llegada" in variables
    assert "sri_ventas_seccion_i" not in variables                       # la trae un extractor


def test_plantilla_llena_desde_el_excel_se_puede_subir():
    libro = load_workbook(io.BytesIO(PLANTILLAS["empleo"].generar_xlsx()))
    libro["Datos"].append(["2026-T2", "EST-01", 12, 3, 5])
    salida = io.BytesIO()
    libro.save(salida)
    r = importar(salida.getvalue(), "plantilla_empleo_v1.xlsx", usuario("analista"), "empleo")
    assert r["estado"] == "publicada"
