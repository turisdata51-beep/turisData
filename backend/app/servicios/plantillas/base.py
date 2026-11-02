"""Marco común de las plantillas de carga (RF-03, RF-05, RF-09).

Cada plantilla declara sus columnas (nombre, tipo, si es obligatoria) y, si hace
falta, reglas de consistencia propias. Con esa misma definición el marco:
- lee el archivo (.xlsx o .csv),
- valida cada fila (tipos, listas, permisos del hotel asociado, duplicados),
- guarda las filas válidas con historial (un reenvío marca el anterior como no
  vigente) y registra los errores de cada fila,
- genera la plantilla Excel que se descarga desde el panel.
Así la plantilla y el validador nunca se desincronizan.
"""
import calendar
import csv
import hashlib
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import text

from ...extractores.base import clave_geografica
from ...nucleo.db import transaccion

TIPOS = {"entero", "decimal", "texto", "lista", "mes", "trimestre", "anio", "fecha", "establecimiento", "zona"}
FILAS_EXCEL = 500


class ErrorArchivo(Exception):
    """El archivo completo no se puede procesar (formato, columnas, vacío)."""


@dataclass(frozen=True)
class Columna:
    nombre: str
    tipo: str
    descripcion: str
    obligatoria: bool = True
    opciones: tuple[str, ...] = ()  # valores permitidos de una lista fija

    def __post_init__(self):
        assert self.tipo in TIPOS, self.tipo


@dataclass
class Contexto:
    """Datos de referencia que la validación consulta una sola vez por archivo."""
    usuario: dict
    establecimientos: dict  # código Cámara -> {id, tipo_alojamiento, categoria, parroquia}
    geografias: dict        # clave geográfica -> id
    listas: dict            # columna -> valores permitidos (fijos o leídos de la base)
    vistos: set = field(default_factory=set)


class Plantilla:
    codigo = ""
    nombre = ""
    descripcion = ""
    unidad = ""             # qué representa una fila, para las instrucciones
    fuente_codigo = ""      # operativo.fuente_datos.codigo
    version = "v1"
    roles = ("administrador", "analista", "asociado")
    columnas: list[Columna] = []
    reglas_texto: list[str] = []
    tabla = ""              # tabla de reportes por establecimiento (si aplica)
    clave_periodo = "periodo"
    solo_referencia: tuple[str, ...] = ()  # columnas que se validan pero no se guardan

    # --- puntos de extensión ---------------------------------------------------------------
    def listas_dinamicas(self, con) -> dict:
        """Listas que vienen de la base (por ejemplo, variables externas)."""
        return {}

    def reglas(self, fila: dict) -> list[tuple[str, str]]:
        """Reglas de consistencia entre columnas: [(campo, mensaje de error)]."""
        return []

    def advertencias(self, fila: dict, contexto: Contexto) -> list[tuple[str, str, str]]:
        """Diferencias que no impiden la carga: [(campo, valor, mensaje)]."""
        return []

    def clave_duplicado(self, fila: dict) -> tuple:
        return fila.get("establecimiento_id"), fila.get(self.clave_periodo)

    def guardar(self, con, filas: list[dict], carga_id: int):
        """Por defecto: tabla por establecimiento y período, con reemplazo e historial."""
        valores = [c.nombre for c in self.columnas
                   if c.tipo not in ("establecimiento", "mes", "trimestre", "anio") and c.nombre not in self.solo_referencia]
        for f in filas:
            con.execute(text(f"""UPDATE {self.tabla} SET vigente = false
                                 WHERE establecimiento_id = :establecimiento_id AND {self.clave_periodo} = :{self.clave_periodo}
                                   AND vigente"""), f)
            columnas = ["carga_id", "establecimiento_id", self.clave_periodo, *valores]
            con.execute(text(f"INSERT INTO {self.tabla} ({', '.join(columnas)}) VALUES ({', '.join(':' + c for c in columnas)})"),
                        {**f, "carga_id": carga_id})

    # --- lectura -----------------------------------------------------------------------------
    def leer(self, contenido: bytes, nombre: str) -> list[dict]:
        if nombre.lower().endswith(".xlsx"):
            libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
            hoja = libro["Datos"] if "Datos" in libro.sheetnames else libro.worksheets[0]
            filas = list(hoja.iter_rows(values_only=True))
        elif nombre.lower().endswith(".csv"):
            texto = contenido.decode("utf-8-sig")
            primera = texto.splitlines()[0] if texto else ""
            filas = list(csv.reader(io.StringIO(texto), delimiter=";" if primera.count(";") > primera.count(",") else ","))
        else:
            raise ErrorArchivo("Formato no soportado: use .xlsx o .csv")
        if not filas:
            raise ErrorArchivo("El archivo está vacío")
        encabezado = [str(c).strip() if c is not None else "" for c in filas[0]]
        faltantes = [c.nombre for c in self.columnas if c.nombre not in encabezado]
        if faltantes:
            raise ErrorArchivo(f"Faltan columnas de la plantilla: {', '.join(faltantes)}")
        registros = []
        for numero, valores in enumerate(filas[1:], start=2):
            registro = {col: valores[i] if i < len(valores) else None for i, col in enumerate(encabezado) if col}
            if all(_vacio(v) for v in registro.values()):
                continue
            registro["_fila"] = numero
            registros.append(registro)
        return registros

    # --- validación --------------------------------------------------------------------------
    def validar(self, registro: dict, ctx: Contexto):
        """Devuelve (fila limpia | None, errores, advertencias, campos obligatorios válidos)."""
        limpio, errores, validos = {}, [], 0
        for col in self.columnas:
            crudo = registro.get(col.nombre)
            if _vacio(crudo):
                limpio[col.nombre] = None
                if col.obligatoria:
                    errores.append((col.nombre, None, "obligatorio", "Campo obligatorio vacío"))
                continue
            valor, error = _convertir(col, crudo, ctx)
            if error:
                errores.append((col.nombre, _texto(crudo), col.tipo, error))
                continue
            if col.tipo == "establecimiento":
                limpio["establecimiento_id"] = valor
            elif col.tipo in ("mes", "trimestre"):
                limpio[self.clave_periodo] = valor
            else:
                limpio[col.nombre] = valor
            validos += col.obligatoria

        if not errores:
            errores += [(campo, None, "consistencia", msg) for campo, msg in self.reglas(limpio)]
        if not errores:
            clave = self.clave_duplicado(limpio)
            if clave in ctx.vistos:
                errores.append((self.columnas[0].nombre, None, "duplicado", "Fila repetida para el mismo período en este archivo"))
        if errores:
            return None, errores, [], validos
        ctx.vistos.add(self.clave_duplicado(limpio))
        advertencias = [(c, v, "referencia", m) for c, v, m in self.advertencias(limpio, ctx)]
        return limpio, [], advertencias, validos

    # --- importación completa ----------------------------------------------------------------
    def importar(self, contenido: bytes, nombre: str, usuario: dict) -> dict:
        from ..indicadores import recalcular  # import local: evita un ciclo de módulos

        huella = hashlib.sha256(contenido).hexdigest()
        with transaccion() as con:
            fuente_id = con.execute(text("SELECT id FROM operativo.fuente_datos WHERE codigo = :c"),
                                    dict(c=self.fuente_codigo)).scalar()
            repetida = con.execute(text("""SELECT id FROM operativo.carga_datos
                                           WHERE hash_sha256 = :h AND estado IN ('publicada', 'con_errores')"""),
                                   dict(h=huella)).scalar()
            carga_id = con.execute(text("""
                INSERT INTO operativo.carga_datos (fuente_id, usuario_id, mecanismo, nombre_archivo, hash_sha256,
                                                   version_plantilla, plantilla, estado)
                VALUES (:f, :u, 'archivo', :n, :h, :v, :p, 'recibida') RETURNING id"""),
                dict(f=fuente_id, u=usuario["id"], n=nombre, h=huella, v=self.version, p=self.codigo)).scalar()

            def rechazar(mensaje):
                con.execute(text("""INSERT INTO operativo.carga_error (carga_id, regla, mensaje, severidad)
                                    VALUES (:c, 'archivo', :m, 'error')"""), dict(c=carga_id, m=mensaje))
                con.execute(text("UPDATE operativo.carga_datos SET estado = 'rechazada' WHERE id = :c"), dict(c=carga_id))
                return _resumen(con, carga_id)

            if repetida:
                return rechazar(f"Este archivo ya se cargó antes (carga #{repetida})")
            try:
                registros = self.leer(contenido, nombre)
            except ErrorArchivo as e:
                return rechazar(str(e))
            except Exception:
                return rechazar("No se pudo leer el archivo; verifique que sea la plantilla oficial")

            ctx = self._contexto(con, usuario)
            validas, problemas, campos_validos = [], [], 0
            for r in registros:
                limpio, errores, advertencias, n = self.validar(r, ctx)
                campos_validos += n
                problemas += [(r["_fila"], *e, "error") for e in errores]
                problemas += [(r["_fila"], *a, "advertencia") for a in advertencias]
                if limpio:
                    validas.append(limpio)

            self.guardar(con, validas, carga_id)
            if problemas:
                con.execute(text("""INSERT INTO operativo.carga_error (carga_id, fila, campo, valor, regla, mensaje, severidad)
                                    VALUES (:c, :fila, :campo, :valor, :regla, :mensaje, :sev)"""),
                            [dict(c=carga_id, fila=f, campo=c, valor=v, regla=rg, mensaje=m, sev=s)
                             for f, c, v, rg, m, s in problemas])
            rechazadas = len(registros) - len(validas)
            obligatorias = sum(c.obligatoria for c in self.columnas)
            con.execute(text("""
                UPDATE operativo.carga_datos SET registros_recibidos = :rec, registros_validos = :val,
                       registros_rechazados = :rech, campos_criticos_esperados = :ce, campos_criticos_validos = :cv,
                       estado = :e WHERE id = :c"""),
                dict(rec=len(registros), val=len(validas), rech=rechazadas, ce=len(registros) * obligatorias,
                     cv=campos_validos, c=carga_id,
                     e="rechazada" if not validas else ("con_errores" if rechazadas else "publicada")))
            resumen = _resumen(con, carga_id)

        if validas:
            recalcular()
        return resumen

    def _contexto(self, con, usuario) -> Contexto:
        establecimientos = {r.codigo_camara: dict(r._mapping) for r in con.execute(text(
            """SELECT e.id, e.codigo_camara, e.tipo_alojamiento, e.categoria, g.nombre AS parroquia
               FROM operativo.establecimiento e JOIN operativo.geografia g ON g.id = e.parroquia_id"""))}
        geografias = {clave_geografica(n): i for n, i in con.execute(text("SELECT nombre, id FROM operativo.geografia")).all()}
        listas = {c.nombre: c.opciones for c in self.columnas if c.opciones} | self.listas_dinamicas(con)
        return Contexto(usuario, establecimientos, geografias, listas)

    # --- plantilla Excel y CSV ---------------------------------------------------------------
    def generar_xlsx(self) -> bytes:
        with transaccion() as con:
            listas = {c.nombre: c.opciones for c in self.columnas if c.opciones} | self.listas_dinamicas(con)
            if any(c.tipo == "zona" for c in self.columnas):
                listas.setdefault("zona", tuple(n for (n,) in con.execute(text(
                    "SELECT nombre FROM operativo.geografia ORDER BY nivel, nombre")).all()))
        libro = Workbook()
        self._hoja_instrucciones(libro.active)
        rangos = self._hoja_listas(libro.create_sheet("Listas"), listas)
        self._hoja_datos(libro.create_sheet("Datos", 1), rangos)
        self._hoja_diccionario(libro.create_sheet("Diccionario"))
        salida = io.BytesIO()
        libro.save(salida)
        return salida.getvalue()

    def generar_csv(self) -> bytes:
        return ("﻿" + ",".join(c.nombre for c in self.columnas) + "\n").encode("utf-8")

    def nombre_archivo(self, extension: str) -> str:
        return f"plantilla_{self.codigo}_{self.version}.{extension}"

    def _hoja_instrucciones(self, ws):
        ws.title = "Instrucciones"
        lineas = [
            (f"{self.nombre} – TurisData, Observatorio Turístico de Cuenca", True),
            (f"Plantilla {self.codigo} {self.version}. No cambie el nombre ni el orden de las columnas.", False),
            (self.descripcion, False),
            ("", False),
            (f"1. Llene una fila por {self.unidad} en la hoja 'Datos'.", False),
            ("2. Las columnas con encabezado rojo son obligatorias; las grises son opcionales.", False),
            ("3. Escriba los números sin separadores de miles (1250, no 1.250).", False),
            ("4. Las cifras se publican solo agregadas, nunca por establecimiento.", False),
            ("5. No incluya nombres, cédulas ni datos de huéspedes o empleados.", False),
            ("6. Pase el cursor sobre cada encabezado para ver su definición.", False),
        ]
        if self.reglas_texto:
            lineas += [("", False), ("Reglas que verifica el sistema al cargar:", True)]
            lineas += [(f"   • {r}", False) for r in self.reglas_texto]
        for i, (texto_linea, negrita) in enumerate(lineas, start=1):
            ws.cell(row=i, column=1, value=texto_linea).font = Font(bold=negrita, size=12 if i == 1 else 11)
        ws.column_dimensions["A"].width = 110

    @staticmethod
    def _hoja_listas(ws, listas: dict) -> dict:
        rangos = {}
        for col, (nombre, valores) in enumerate(listas.items(), start=1):
            letra = ws.cell(row=1, column=col).column_letter
            ws.cell(row=1, column=col, value=nombre).font = Font(bold=True)
            for fila, valor in enumerate(valores, start=2):
                ws.cell(row=fila, column=col, value=valor)
            ws.column_dimensions[letra].width = 30
            rangos[nombre] = f"Listas!${letra}$2:${letra}${len(valores) + 1}"
        return rangos

    def _hoja_datos(self, ws, rangos: dict):
        obligatoria, opcional = PatternFill("solid", fgColor="A5162D"), PatternFill("solid", fgColor="7F7F7F")
        ultima = FILAS_EXCEL + 1
        for i, col in enumerate(self.columnas, start=1):
            celda = ws.cell(row=1, column=i, value=col.nombre)
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = obligatoria if col.obligatoria else opcional
            celda.alignment = Alignment(wrap_text=True, vertical="center")
            celda.comment = Comment(col.descripcion, "TurisData")
            letra = celda.column_letter
            ws.column_dimensions[letra].width = max(16, len(col.nombre) + 2)
            rango = f"{letra}2:{letra}{ultima}"
            validacion = None
            if col.nombre in rangos:
                validacion = DataValidation(type="list", formula1=f"={rangos[col.nombre]}", allow_blank=True,
                                            error="Elija un valor de la lista.", errorTitle="Valor no permitido")
            elif col.tipo == "entero":
                validacion = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                                            error="Ingrese un número entero mayor o igual a 0.", errorTitle="Número inválido")
            elif col.tipo == "decimal":
                validacion = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                                            error="Ingrese un valor mayor o igual a 0.", errorTitle="Número inválido")
            elif col.tipo in ("mes", "trimestre", "anio", "fecha"):
                for fila in range(2, ultima + 1):
                    ws.cell(row=fila, column=i).number_format = "@"  # que Excel no lo convierta en fecha
            if validacion:
                ws.add_data_validation(validacion)
                validacion.add(rango)
        ws.row_dimensions[1].height = 45
        ws.freeze_panes = "B2"

    def _hoja_diccionario(self, ws):
        for col, titulo in enumerate(["columna", "obligatoria", "tipo", "descripción"], start=1):
            ws.cell(row=1, column=col, value=titulo).font = Font(bold=True)
        for fila, c in enumerate(self.columnas, start=2):
            for col, valor in enumerate([c.nombre, "Sí" if c.obligatoria else "No", c.tipo, c.descripcion], start=1):
                ws.cell(row=fila, column=col, value=valor)
        for letra, ancho in zip("ABCD", (32, 12, 14, 95)):
            ws.column_dimensions[letra].width = ancho


# --- conversión de valores ----------------------------------------------------------------------

def _vacio(valor) -> bool:
    return valor is None or str(valor).strip() == ""


def _texto(valor) -> str:
    return valor.strftime("%Y-%m-%d") if hasattr(valor, "strftime") else str(valor).strip()


def _convertir(col: Columna, crudo, ctx: Contexto):
    """(valor convertido, None) o (None, mensaje de error)."""
    valor = _texto(crudo)
    if col.tipo == "entero":
        try:
            numero = Decimal(valor)
            if numero != numero.to_integral_value() or numero < 0:
                raise InvalidOperation
            return int(numero), None
        except InvalidOperation:
            return None, "Debe ser un número entero mayor o igual a 0"
    if col.tipo == "decimal":
        try:
            numero = Decimal(valor.replace(",", "."))
            if numero < 0:
                raise InvalidOperation
            return numero, None
        except InvalidOperation:
            return None, "Debe ser un valor mayor o igual a 0"
    if col.tipo == "mes":
        texto = crudo.strftime("%Y-%m") if isinstance(crudo, (date, datetime)) else valor[:7]
        if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", texto):
            return None, "Use el formato AAAA-MM"
        return date(int(texto[:4]), int(texto[5:7]), 1), None
    if col.tipo == "trimestre":
        m = re.fullmatch(r"(\d{4})-T([1-4])", valor.upper())
        if not m:
            return None, "Use el formato AAAA-T1 a AAAA-T4"
        return date(int(m.group(1)), (int(m.group(2)) - 1) * 3 + 1, 1), None
    if col.tipo == "anio":
        if not re.fullmatch(r"(19|20)\d{2}", valor.split(".")[0]):
            return None, "Use un año de cuatro dígitos"
        return int(valor.split(".")[0]), None
    if col.tipo == "fecha":
        try:
            return (crudo.date() if isinstance(crudo, datetime) else crudo if isinstance(crudo, date)
                    else date.fromisoformat(valor[:10])), None
        except ValueError:
            return None, "Use el formato AAAA-MM-DD"
    if col.tipo == "lista":
        permitidos = ctx.listas.get(col.nombre, ())
        if permitidos and valor not in permitidos:
            return None, "Valor fuera de la lista permitida"
        return valor, None
    if col.tipo == "establecimiento":
        est = ctx.establecimientos.get(valor)
        if not est:
            return None, "Código no registrado en la Cámara"
        if ctx.usuario["rol"] == "asociado" and est["id"] != ctx.usuario["establecimiento_id"]:
            return None, "Solo puede reportar su propio establecimiento"
        return est["id"], None
    if col.tipo == "zona":
        geografia = ctx.geografias.get(clave_geografica(valor))
        return (geografia, None) if geografia else (None, "Zona no reconocida (use Cuenca o el nombre de una parroquia)")
    return valor, None  # texto


def dias_del_mes(periodo: date) -> int:
    return calendar.monthrange(periodo.year, periodo.month)[1]


def _resumen(con, carga_id) -> dict:
    carga = dict(con.execute(text("SELECT * FROM operativo.carga_datos WHERE id = :c"), dict(c=carga_id)).mappings().one())
    carga["errores"] = [dict(e) for e in con.execute(text(
        """SELECT fila, campo, valor, regla, mensaje, severidad FROM operativo.carga_error
           WHERE carga_id = :c ORDER BY fila NULLS FIRST, id"""), dict(c=carga_id)).mappings()]
    return carga
