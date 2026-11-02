"""INEC: estimaciones y proyecciones de población cantonal (Revisión 2024).

Descarga el tabulado XLSX y lee la hoja 3.2 (cantones por grupos de edad y sexo).
Toma la fila "Total" de Cuenca y la columna "Total" de cada año (2010–2035).
Las cifras están corregidas por omisión censal, por eso 2022 no coincide con el
conteo del censo (596.101); los indicadores las marcan como estimadas.
"""
import io
import json
from datetime import date

from openpyxl import load_workbook

from .base import Dato, Extractor

HOJA = "3.2"
CANTON = "Cuenca"


class InecPoblacion(Extractor):
    fuente_codigo = "inec_poblacion"
    url = ("https://www.ecuadorencifras.gob.ec/documentos/web-inec/Poblacion_y_Demografia/Proyecciones_Poblacionales/"
           "censo_2022/revision_2024_areas/Tabulados_Estimacion_Proyeccion_v2.xlsx")

    def descargar(self) -> bytes:
        libro = load_workbook(io.BytesIO(self.obtener(self.url).content), read_only=True)
        if HOJA not in libro.sheetnames:
            raise ValueError(f"El tabulado ya no tiene la hoja {HOJA}; revise si cambió el formato")
        filas = libro[HOJA].iter_rows(values_only=True)
        columna_anio = {}
        for fila in filas:
            # La fila de años tiene un año cada tres columnas (Hombre, Mujer, Total)
            anios = {j: str(v).strip()[:4] for j, v in enumerate(fila) if v is not None and str(v).strip()[:4].isdigit()}
            if len(anios) > 10:
                columna_anio = {int(a): j for j, a in anios.items()}
                break
        poblacion = {}
        for fila in filas:
            textos = [str(v).strip() if v is not None else "" for v in fila]
            if CANTON in textos and "Total" in textos:
                poblacion = {anio: int(fila[j + 2]) for anio, j in columna_anio.items()}
                break
        if not poblacion:
            raise ValueError("No se encontró el total de Cuenca en el tabulado")
        return json.dumps(poblacion, sort_keys=True).encode("utf-8")

    def interpretar(self, contenido: bytes) -> list[Dato]:
        return [Dato("inec_poblacion_proyectada", date(int(anio), 1, 1), date(int(anio), 12, 31), habitantes,
                     periodicidad="anual")
                for anio, habitantes in json.loads(contenido).items()]
