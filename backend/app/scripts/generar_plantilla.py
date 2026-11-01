"""Genera la plantilla de reporte mensual de alojamiento (XLSX + CSV).

Uso (carpeta backend):  python -m app.scripts.generar_plantilla
Salida: recursos/plantillas/reporte_mensual_alojamiento_v1.xlsx y .csv

Las columnas deben coincidir con recursos/diccionario_indicadores.yaml (variables_base)
y con la tabla operativo.reporte_alojamiento_mensual de db/001_esquema.sql.
"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

VERSION = "v1"
CARPETA = Path(__file__).resolve().parents[2] / "recursos" / "plantillas"

# (columna, obligatorio, tipo, descripción)
COLUMNAS = [
    ("periodo", True, "texto", "Mes reportado en formato AAAA-MM (ej. 2026-09)."),
    ("codigo_establecimiento", True, "texto", "Código asignado por la Cámara (no usar RUC ni cédula)."),
    ("tipo_alojamiento", True, "lista", "Clasificación según el Reglamento de Alojamiento Turístico."),
    ("categoria", True, "lista", "Categoría registrada en el Catastro de MINTUR."),
    ("parroquia", True, "lista", "Parroquia del cantón Cuenca donde está el establecimiento."),
    ("habitaciones_instaladas", True, "entero", "Habitaciones que tiene el establecimiento (capacidad física)."),
    ("plazas_instaladas", True, "entero", "Camas/plazas instaladas. Una cama doble = 2 plazas."),
    ("dias_operacion", True, "entero", "Días del mes en que el establecimiento estuvo abierto."),
    ("habitaciones_noche_disponibles", True, "entero",
     "Suma, día por día, de las habitaciones a la venta (sin las fuera de servicio)."),
    ("habitaciones_noche_ocupadas", True, "entero", "Suma, día por día, de las habitaciones vendidas."),
    ("plazas_cama_noche_disponibles", True, "entero", "Suma, día por día, de las plazas a la venta."),
    ("llegadas_residentes", True, "entero", "Huéspedes que hicieron check-in en el mes y viven en Ecuador."),
    ("llegadas_no_residentes", True, "entero", "Huéspedes que hicieron check-in en el mes y viven fuera de Ecuador."),
    ("noches_residentes", True, "entero", "Noches pasadas en el mes por huéspedes residentes (1 persona x 1 noche = 1)."),
    ("noches_no_residentes", True, "entero", "Noches pasadas en el mes por huéspedes no residentes."),
    ("ingreso_habitaciones_usd", False, "decimal",
     "OPCIONAL. Venta de habitaciones del mes en USD, sin IVA ni 10 % de servicio."),
    ("observaciones", False, "texto", "OPCIONAL. Cierres, remodelaciones, eventos u otras novedades."),
]

TIPOS_ALOJAMIENTO = [
    "Hotel", "Hostal", "Hostería", "Hacienda turística", "Lodge", "Resort",
    "Refugio", "Campamento turístico", "Casa de huéspedes",
]
CATEGORIAS = ["1 estrella", "2 estrellas", "3 estrellas", "4 estrellas", "5 estrellas", "Categoría única"]
PARROQUIAS_URBANAS = [
    "Bellavista", "Cañaribamba", "El Batán", "El Sagrario", "El Vecino", "Gil Ramírez Dávalos",
    "Hermano Miguel", "Huayna Cápac", "Machángara", "Monay", "San Blas", "San Sebastián",
    "Sucre", "Totoracocha", "Yanuncay",
]
PARROQUIAS_RURALES = [
    "Baños", "Chaucha", "Checa", "Chiquintad", "Cumbe", "El Valle", "Llacao", "Molleturo",
    "Nulti", "Octavio Cordero Palacios", "Paccha", "Quingeo", "Ricaurte", "San Joaquín",
    "Santa Ana", "Sayausí", "Sidcay", "Sinincay", "Tarqui", "Turi", "Victoria del Portete",
]

# Reglas de consistencia que aplicará el validador del sistema (se documentan en la hoja).
REGLAS = [
    "habitaciones_noche_ocupadas <= habitaciones_noche_disponibles",
    "habitaciones_noche_disponibles <= habitaciones_instaladas x días del mes",
    "plazas_cama_noche_disponibles <= plazas_instaladas x días del mes",
    "plazas_cama_noche_disponibles >= habitaciones_noche_disponibles",
    "noches_residentes + noches_no_residentes <= plazas_cama_noche_disponibles",
    "dias_operacion entre 0 y los días del mes",
    "Si hay habitaciones ocupadas debe haber noches > 0, y viceversa",
    "Un solo reporte vigente por establecimiento y mes (un reenvío reemplaza al anterior)",
]

FILAS_DATOS = 500
ENCABEZADO = PatternFill("solid", fgColor="1F4E78")
OPCIONAL = PatternFill("solid", fgColor="7F7F7F")


def hoja_instrucciones(wb):
    ws = wb.active
    ws.title = "Instrucciones"
    lineas = [
        ("Reporte mensual de alojamiento – Observatorio Turístico Inteligente de Cuenca", True),
        (f"Plantilla {VERSION}. No cambie el nombre ni el orden de las columnas.", False),
        ("", False),
        ("1. Llene una fila por establecimiento y mes en la hoja 'Datos'.", False),
        ("2. Los campos con encabezado azul son obligatorios; los grises son opcionales.", False),
        ("3. Use solo números enteros sin separadores de miles (ej. 1250, no 1.250).", False),
        ("4. Los datos se publican solo agregados y nunca por establecimiento individual.", False),
        ("5. No incluya nombres, cédulas ni datos de huéspedes.", False),
        ("6. Pase el cursor sobre cada encabezado para ver su definición.", False),
        ("", False),
        ("Reglas de consistencia que verifica el sistema al cargar:", True),
    ] + [(f"   • {r}", False) for r in REGLAS]
    for i, (texto, negrita) in enumerate(lineas, start=1):
        ws.cell(row=i, column=1, value=texto).font = Font(bold=negrita, size=12 if i == 1 else 11)
    ws.column_dimensions["A"].width = 110


def hoja_catalogos(wb):
    ws = wb.create_sheet("Catalogos")
    listas = {"tipo_alojamiento": TIPOS_ALOJAMIENTO, "categoria": CATEGORIAS,
              "parroquia": PARROQUIAS_URBANAS + PARROQUIAS_RURALES}
    rangos = {}
    for col, (nombre, valores) in enumerate(listas.items(), start=1):
        letra = ws.cell(row=1, column=col).column_letter
        ws.cell(row=1, column=col, value=nombre).font = Font(bold=True)
        for fila, valor in enumerate(valores, start=2):
            ws.cell(row=fila, column=col, value=valor)
        ws.column_dimensions[letra].width = 28
        rangos[nombre] = f"Catalogos!${letra}$2:${letra}${len(valores) + 1}"
    return rangos


def hoja_diccionario(wb):
    ws = wb.create_sheet("Diccionario")
    for col, titulo in enumerate(["columna", "obligatorio", "tipo", "descripción"], start=1):
        ws.cell(row=1, column=col, value=titulo).font = Font(bold=True)
    for fila, (nombre, oblig, tipo, desc) in enumerate(COLUMNAS, start=2):
        ws.cell(row=fila, column=1, value=nombre)
        ws.cell(row=fila, column=2, value="Sí" if oblig else "No")
        ws.cell(row=fila, column=3, value=tipo)
        ws.cell(row=fila, column=4, value=desc)
    for letra, ancho in zip("ABCD", (32, 12, 10, 95)):
        ws.column_dimensions[letra].width = ancho


def hoja_datos(wb, rangos):
    ws = wb.create_sheet("Datos", 1)
    for col, (nombre, oblig, tipo, desc) in enumerate(COLUMNAS, start=1):
        celda = ws.cell(row=1, column=col, value=nombre)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = ENCABEZADO if oblig else OPCIONAL
        celda.alignment = Alignment(wrap_text=True, vertical="center")
        celda.comment = Comment(desc, "Observatorio")
        ws.column_dimensions[celda.column_letter].width = max(16, len(nombre) + 2)
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = "C2"

    ultima = FILAS_DATOS + 1
    letra = {nombre: ws.cell(row=1, column=i).column_letter for i, (nombre, *_) in enumerate(COLUMNAS, 1)}

    dv_periodo = DataValidation(
        type="custom", allow_blank=True,
        formula1='AND(LEN(A2)=7,MID(A2,5,1)="-",ISNUMBER(--LEFT(A2,4)),'
                 'ISNUMBER(--RIGHT(A2,2)),--RIGHT(A2,2)>=1,--RIGHT(A2,2)<=12)',
        error="Use el formato AAAA-MM, por ejemplo 2026-09.", errorTitle="Período inválido")
    ws.add_data_validation(dv_periodo)
    dv_periodo.add(f"A2:A{ultima}")
    for fila in range(2, ultima + 1):
        ws.cell(row=fila, column=1).number_format = "@"

    for nombre in ("tipo_alojamiento", "categoria", "parroquia"):
        dv = DataValidation(type="list", formula1=f"={rangos[nombre]}", allow_blank=True,
                            error="Elija un valor de la lista.", errorTitle="Valor no permitido")
        ws.add_data_validation(dv)
        dv.add(f"{letra[nombre]}2:{letra[nombre]}{ultima}")

    for nombre, _, tipo, _ in COLUMNAS:
        if tipo == "entero":
            dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                                error="Ingrese un número entero mayor o igual a 0.", errorTitle="Número inválido")
        elif tipo == "decimal":
            dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                                error="Ingrese un valor mayor o igual a 0.", errorTitle="Número inválido")
        else:
            continue
        ws.add_data_validation(dv)
        dv.add(f"{letra[nombre]}2:{letra[nombre]}{ultima}")

    # Regla cruzada visible en Excel: ocupadas <= disponibles
    occ, disp = letra["habitaciones_noche_ocupadas"], letra["habitaciones_noche_disponibles"]
    dv_occ = DataValidation(type="custom", allow_blank=True, formula1=f"{occ}2<={disp}2",
                            error="Las habitaciones ocupadas no pueden superar las disponibles.",
                            errorTitle="Inconsistencia")
    ws.add_data_validation(dv_occ)
    dv_occ.add(f"{occ}2:{occ}{ultima}")
    return ws


def main():
    wb = Workbook()
    hoja_instrucciones(wb)
    rangos = hoja_catalogos(wb)
    hoja_datos(wb, rangos)
    hoja_diccionario(wb)
    xlsx = CARPETA / f"reporte_mensual_alojamiento_{VERSION}.xlsx"
    wb.save(xlsx)

    csv_path = CARPETA / f"reporte_mensual_alojamiento_{VERSION}.csv"
    with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
        csv.writer(f).writerow([c[0] for c in COLUMNAS])
    print(f"Generados: {xlsx.name}, {csv_path.name}")


if __name__ == "__main__":
    main()
