"""Importa la plantilla mensual de alojamiento (XLSX o CSV) y registra la trazabilidad (RF-03, RF-05, RF-09).

Las filas válidas se guardan; las inválidas quedan en operativo.carga_error.
Un reporte nuevo del mismo establecimiento y mes reemplaza al anterior (vigente = false).
"""
import calendar
import csv
import hashlib
import io
import re
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from sqlalchemy import text

from .calculo import recalcular
from .db import transaccion

VERSION_PLANTILLA = "v1"
ENTEROS = [
    "habitaciones_instaladas", "plazas_instaladas", "dias_operacion",
    "habitaciones_noche_disponibles", "habitaciones_noche_ocupadas", "plazas_cama_noche_disponibles",
    "llegadas_residentes", "llegadas_no_residentes", "noches_residentes", "noches_no_residentes",
]
TEXTOS_OBLIGATORIOS = ["periodo", "codigo_establecimiento", "tipo_alojamiento", "categoria", "parroquia"]
COLUMNAS = TEXTOS_OBLIGATORIOS + ENTEROS + ["ingreso_habitaciones_usd", "observaciones"]
CRITICOS = TEXTOS_OBLIGATORIOS + ENTEROS


class ErrorCarga(Exception):
    pass


def _leer(contenido: bytes, nombre: str) -> list[dict]:
    if nombre.lower().endswith(".xlsx"):
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        hoja = libro["Datos"] if "Datos" in libro.sheetnames else libro.worksheets[0]
        filas = list(hoja.iter_rows(values_only=True))
    elif nombre.lower().endswith(".csv"):
        texto = contenido.decode("utf-8-sig")
        separador = ";" if texto.splitlines()[0].count(";") > texto.splitlines()[0].count(",") else ","
        filas = list(csv.reader(io.StringIO(texto), delimiter=separador))
    else:
        raise ErrorCarga("Formato no soportado: use .xlsx o .csv")
    if not filas:
        raise ErrorCarga("El archivo está vacío")
    encabezado = [str(c).strip() if c is not None else "" for c in filas[0]]
    faltantes = [c for c in COLUMNAS if c not in encabezado]
    if faltantes:
        raise ErrorCarga(f"Faltan columnas de la plantilla: {', '.join(faltantes)}")
    registros = []
    for numero, valores in enumerate(filas[1:], start=2):
        registro = {col: valores[i] if i < len(valores) else None for i, col in enumerate(encabezado) if col}
        if all(v is None or str(v).strip() == "" for v in registro.values()):
            continue
        registro["_fila"] = numero
        registros.append(registro)
    return registros


def _vacio(valor) -> bool:
    return valor is None or str(valor).strip() == ""


def _validar(r: dict, establecimientos: dict, usuario: dict, vistos: set):
    """Devuelve (datos_limpios | None, errores, advertencias, campos_criticos_validos)."""
    errores, advertencias, limpio = [], [], {}
    validos = 0

    for campo in TEXTOS_OBLIGATORIOS:
        if _vacio(r.get(campo)):
            errores.append((campo, None, "obligatorio", "Campo obligatorio vacío"))
        else:
            validos += 1

    periodo = r.get("periodo")
    # Excel a veces convierte "2026-09" en fecha
    periodo = periodo.strftime("%Y-%m") if hasattr(periodo, "strftime") else str(periodo or "").strip()[:7]
    if periodo and not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", periodo):
        errores.append(("periodo", periodo, "formato", "Use el formato AAAA-MM"))
        validos -= 1
    codigo = str(r.get("codigo_establecimiento") or "").strip()
    est = establecimientos.get(codigo)
    if codigo and not est:
        errores.append(("codigo_establecimiento", codigo, "existe", "Código no registrado en la Cámara"))
        validos -= 1
    elif est and usuario["rol"] == "asociado" and est["id"] != usuario["establecimiento_id"]:
        errores.append(("codigo_establecimiento", codigo, "permiso", "Solo puede reportar su propio establecimiento"))

    for campo in ENTEROS:
        valor = r.get(campo)
        if _vacio(valor):
            errores.append((campo, None, "obligatorio", "Campo obligatorio vacío"))
            continue
        try:
            numero = Decimal(str(valor).strip())
            if numero != numero.to_integral_value() or numero < 0:
                raise InvalidOperation
            limpio[campo] = int(numero)
            validos += 1
        except InvalidOperation:
            errores.append((campo, str(valor), "entero", "Debe ser un número entero mayor o igual a 0"))

    ingreso = r.get("ingreso_habitaciones_usd")
    limpio["ingreso_habitaciones_usd"] = None
    if not _vacio(ingreso):
        try:
            limpio["ingreso_habitaciones_usd"] = Decimal(str(ingreso).strip().replace(",", "."))
            if limpio["ingreso_habitaciones_usd"] < 0:
                raise InvalidOperation
        except InvalidOperation:
            errores.append(("ingreso_habitaciones_usd", str(ingreso), "decimal", "Debe ser un valor mayor o igual a 0"))

    formato_ok = re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", periodo) and all(c in limpio for c in ENTEROS)
    if formato_ok:
        anio, mes = map(int, periodo.split("-"))
        dias = calendar.monthrange(anio, mes)[1]
        d = limpio
        noches = d["noches_residentes"] + d["noches_no_residentes"]
        reglas = [
            (d["habitaciones_noche_ocupadas"] <= d["habitaciones_noche_disponibles"],
             "habitaciones_noche_ocupadas", "Las habitaciones ocupadas superan las disponibles"),
            (d["dias_operacion"] <= dias, "dias_operacion", f"El mes tiene {dias} días"),
            (d["habitaciones_instaladas"] > 0 and d["plazas_instaladas"] > 0,
             "habitaciones_instaladas", "La capacidad instalada debe ser mayor a 0"),
            (d["habitaciones_noche_disponibles"] <= d["habitaciones_instaladas"] * dias,
             "habitaciones_noche_disponibles", "Supera habitaciones instaladas × días del mes"),
            (d["plazas_cama_noche_disponibles"] <= d["plazas_instaladas"] * dias,
             "plazas_cama_noche_disponibles", "Supera plazas instaladas × días del mes"),
            (d["plazas_cama_noche_disponibles"] >= d["habitaciones_noche_disponibles"],
             "plazas_cama_noche_disponibles", "Las plazas no pueden ser menos que las habitaciones"),
            (noches <= d["plazas_cama_noche_disponibles"],
             "noches_residentes", "Las noches superan las plazas disponibles"),
            ((d["habitaciones_noche_ocupadas"] == 0) == (noches == 0),
             "habitaciones_noche_ocupadas", "Si hay habitaciones ocupadas debe haber noches, y viceversa"),
        ]
        errores += [(campo, None, "consistencia", msg) for ok, campo, msg in reglas if not ok]
        if (codigo, periodo) in vistos:
            errores.append(("codigo_establecimiento", codigo, "duplicado", "Fila repetida para el mismo mes en este archivo"))

    if est and not errores:
        for campo in ("tipo_alojamiento", "categoria"):
            if str(r.get(campo)).strip() != est[campo]:
                advertencias.append((campo, str(r.get(campo)), "maestro",
                                     f"No coincide con el registro del establecimiento ({est[campo]})"))
        if str(r.get("parroquia")).strip() != est["parroquia"]:
            advertencias.append(("parroquia", str(r.get("parroquia")), "maestro",
                                 f"No coincide con el registro del establecimiento ({est['parroquia']})"))

    if errores:
        return None, errores, advertencias, max(validos, 0)
    vistos.add((codigo, periodo))
    limpio.update(establecimiento_id=est["id"], periodo=f"{periodo}-01",
                  observaciones=None if _vacio(r.get("observaciones")) else str(r["observaciones"]).strip())
    return limpio, [], advertencias, validos


def importar(contenido: bytes, nombre: str, usuario: dict) -> dict:
    huella = hashlib.sha256(contenido).hexdigest()
    with transaccion() as con:
        fuente_id = con.execute(text(
            "SELECT id FROM operativo.fuente_datos WHERE codigo = 'plantilla_alojamiento_camara'")).scalar()
        repetida = con.execute(text(
            """SELECT id FROM operativo.carga_datos
               WHERE hash_sha256 = :h AND estado IN ('publicada', 'con_errores')"""), dict(h=huella)).scalar()
        carga_id = con.execute(text(
            """INSERT INTO operativo.carga_datos (fuente_id, usuario_id, mecanismo, nombre_archivo, hash_sha256,
                                                   version_plantilla, estado)
               VALUES (:f, :u, 'archivo', :n, :h, :v, 'recibida') RETURNING id"""),
            dict(f=fuente_id, u=usuario["id"], n=nombre, h=huella, v=VERSION_PLANTILLA)).scalar()

        def error_general(mensaje):
            con.execute(text(
                """INSERT INTO operativo.carga_error (carga_id, regla, mensaje, severidad)
                   VALUES (:c, 'archivo', :m, 'error')"""), dict(c=carga_id, m=mensaje))
            con.execute(text("UPDATE operativo.carga_datos SET estado = 'rechazada' WHERE id = :c"), dict(c=carga_id))

        if repetida:
            error_general(f"Este archivo ya se cargó antes (carga #{repetida})")
            return _resumen(con, carga_id)
        try:
            registros = _leer(contenido, nombre)
        except ErrorCarga as e:
            error_general(str(e))
            return _resumen(con, carga_id)
        except Exception:
            error_general("No se pudo leer el archivo; verifique que sea la plantilla oficial")
            return _resumen(con, carga_id)

        establecimientos = {r.codigo_camara: dict(r._mapping) for r in con.execute(text(
            """SELECT e.id, e.codigo_camara, e.tipo_alojamiento, e.categoria, g.nombre AS parroquia
               FROM operativo.establecimiento e JOIN operativo.geografia g ON g.id = e.parroquia_id"""))}
        vistos, validos, campos_validos, problemas = set(), [], 0, []
        for r in registros:
            limpio, errores, advertencias, n_campos = _validar(r, establecimientos, usuario, vistos)
            campos_validos += n_campos
            problemas += [(r["_fila"], *e, "error") for e in errores]
            problemas += [(r["_fila"], *a, "advertencia") for a in advertencias]
            if limpio:
                validos.append(limpio)

        for d in validos:
            con.execute(text(
                """UPDATE operativo.reporte_alojamiento_mensual SET vigente = false
                   WHERE establecimiento_id = :establecimiento_id AND periodo = :periodo AND vigente"""), d)
            con.execute(text(
                f"""INSERT INTO operativo.reporte_alojamiento_mensual
                      (carga_id, establecimiento_id, periodo, {', '.join(ENTEROS)}, ingreso_habitaciones_usd, observaciones)
                    VALUES (:carga_id, :establecimiento_id, :periodo, {', '.join(':' + c for c in ENTEROS)},
                            :ingreso_habitaciones_usd, :observaciones)"""), {**d, "carga_id": carga_id})
        if problemas:
            con.execute(text(
                """INSERT INTO operativo.carga_error (carga_id, fila, campo, valor, regla, mensaje, severidad)
                   VALUES (:c, :fila, :campo, :valor, :regla, :mensaje, :sev)"""),
                [dict(c=carga_id, fila=f, campo=c, valor=v, regla=rg, mensaje=m, sev=s)
                 for f, c, v, rg, m, s in problemas])

        rechazados = len(registros) - len(validos)
        estado = "rechazada" if not validos else ("con_errores" if rechazados else "publicada")
        con.execute(text(
            """UPDATE operativo.carga_datos SET registros_recibidos = :rec, registros_validos = :val,
                      registros_rechazados = :rech, campos_criticos_esperados = :ce,
                      campos_criticos_validos = :cv, estado = :e
               WHERE id = :c"""),
            dict(rec=len(registros), val=len(validos), rech=rechazados, ce=len(registros) * len(CRITICOS),
                 cv=campos_validos, e=estado, c=carga_id))
        resumen = _resumen(con, carga_id)

    if validos:
        recalcular()
    return resumen


def _resumen(con, carga_id) -> dict:
    carga = dict(con.execute(text("SELECT * FROM operativo.carga_datos WHERE id = :c"), dict(c=carga_id)).mappings().one())
    carga["errores"] = [dict(e) for e in con.execute(text(
        "SELECT fila, campo, valor, regla, mensaje, severidad FROM operativo.carga_error WHERE carga_id = :c ORDER BY fila NULLS FIRST, id"),
        dict(c=carga_id)).mappings()]
    return carga
